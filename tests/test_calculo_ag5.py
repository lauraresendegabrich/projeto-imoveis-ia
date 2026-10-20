"""
Testes do calculo do Agente 5 (agents/price_liquidity.py) com exemplos reais:
a planilha do professor, um apartamento de BH e uma casa de Curitiba. Sem rede.

    .venv/Scripts/python.exe -m tests.test_calculo_ag5
"""
import sys

sys.path.insert(0, ".")

from agents.price_liquidity import executar_agente5


def apto(preco, area):
    return {"propertyType": "apartamento", "price": preco, "area": area}


def casa(preco, area, area_terreno):
    return {"propertyType": "casa", "price": preco, "area": area, "area_terreno": area_terreno}


def terreno(preco, area):
    return {"propertyType": "terreno", "price": preco, "area": area}


def calcular(alvo, comparaveis, terrenos=()):
    return executar_agente5(alvo, list(terrenos), comparaveis, {}, {})


def valor_medio(r):
    return r["avaliacao_planilha"]["valor_medio_imovel"]


def test_planilha_do_professor_5_aptos():
    # Ficha da Construcao, linhas 38-42; area do imovel avaliado (E7) = 250 m2.
    # Com 5 aptos, contar uma ou duas vezes da o mesmo resultado (C64 e C65).
    comps = [apto(350000, 242), apto(430000, 232), apto(680000, 169), apto(400000, 300), apto(520000, 140)]
    r = calcular({"propertyType": "apartamento", "area": 250}, comps)
    assert valor_medio(r) == 584501.25, valor_medio(r)
    assert r["avaliacao_planilha"]["valor_minimo_imovel"] == 333333.33
    assert r["avaliacao_planilha"]["valor_liquidez"] == 526051.12


def test_bh_floresta_6_aptos_cada_um_conta_uma_vez():
    # Antes (lista dobrada): R$ 5.691,69/m2 -> R$ 620.394,07. Agora: R$ 5.721,10/m2.
    comps = [apto(300000, 93), apto(499900, 120), apto(630000, 110),
             apto(648000, 117), apto(650000, 77), apto(790000, 106)]
    r = calcular({"propertyType": "apartamento", "area": 109}, comps)
    assert r["calculo_construcao"]["valor_m2_referencia"] == 5721.1, r["calculo_construcao"]
    assert valor_medio(r) == 623599.84, valor_medio(r)
    assert r["valor_m2_zona_homogenea"]["construcao"]["combinados"]["quantidade_total"] == 6
    assert r["status_avaliacao"] == "ok"


def test_casa_curitiba_com_4_terrenos_nao_separa_o_terreno():
    # Caso real (Cajuru, 10/10/2026). Antes, com 4 terrenos o terreno era separado
    # (R$ 913.061,64); desde 17/10/2026 so com 5+: calcula pelo m2 das casas.
    comps = [casa(690000, 146, 105), casa(699900, 126, 126), casa(799000, 136, 136), casa(899000, 173, 173)]
    terrenos = [terreno(515000, 540), terreno(789000, 600), terreno(836000, 418), terreno(1950000, 1824)]
    r = calcular({"propertyType": "casa", "area": 171.19, "area_terreno": 146.76}, comps, terrenos)
    assert r["calculo_terreno"]["aplicado"] is False
    assert valor_medio(r) == 920256.98, valor_medio(r)
    assert r["comparaveis_usados"]["terreno"] == [] and len(r["comparaveis_descartados"]["terreno"]) == 4
    assert r["aviso_terreno"].startswith("Só 4 terreno(s) à venda na região (mínimo de 5).")


def test_minimo_de_5_comparaveis():
    alvo = {"propertyType": "apartamento", "area": 100}
    quatro = [apto(500000 + i * 10000, 100) for i in range(4)]
    assert calcular(alvo, quatro)["status_avaliacao"] == "amostra_insuficiente"
    assert calcular(alvo, quatro + [apto(450000, 100)])["status_avaliacao"] == "ok"


CASAS_5 = [casa(600000 + i * 50000, 150, 200) for i in range(5)]
ALVO_CASA = {"propertyType": "casa", "area": 150, "area_terreno": 200}


def test_casa_sem_nenhum_terreno_usa_preco_m2_total():
    r = calcular(ALVO_CASA, CASAS_5, [])
    assert r["status_avaliacao"] == "ok", r["status_avaliacao"]
    assert r["calculo_terreno"]["aplicado"] is False
    assert r["aviso_terreno"] is None


def test_casa_com_3_terrenos_calcula_sem_separar_com_aviso():
    # Decisao da usuaria (17/10/2026): com 1 a 4 terrenos, calcula como sem terreno.
    sem_terreno = calcular(ALVO_CASA, CASAS_5, [])
    terrenos = [terreno(200000 + i * 10000, 250) for i in range(3)]
    r = calcular(ALVO_CASA, CASAS_5, terrenos)
    assert r["status_avaliacao"] == "ok", r["status_avaliacao"]
    assert r["calculo_terreno"]["aplicado"] is False
    assert valor_medio(r) == valor_medio(sem_terreno)
    assert r["comparaveis_usados"]["terreno"] == [] and len(r["comparaveis_descartados"]["terreno"]) == 3
    assert r["aviso_terreno"] == (
        "Só 3 terreno(s) à venda na região (mínimo de 5). O valor foi calculado pelo m² das casas "
        "parecidas, sem separar o terreno; se o lote do imóvel for bem maior ou menor que o das "
        "casas comparadas, o valor fica menos preciso."), r["aviso_terreno"]


def test_casa_com_5_terrenos_separa_terreno():
    terrenos = [terreno(200000 + i * 10000, 250) for i in range(5)]
    r = calcular(ALVO_CASA, CASAS_5, terrenos)
    assert r["status_avaliacao"] == "ok" and r["calculo_terreno"]["aplicado"] is True, r["status_avaliacao"]
    assert len(r["comparaveis_usados"]["terreno"]) == 5 and r["aviso_terreno"] is None
    # Casa separando o terreno: as duas series (menor e medio m2 de terreno) continuam juntas.
    assert r["valor_m2_zona_homogenea"]["construcao"]["combinados"]["quantidade_total"] == 10



def test_banco_e_excel_com_terreno_nao_separado():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado, montar_registro
    r = calcular(ALVO_CASA, CASAS_5, [terreno(200000 + i * 10000, 250) for i in range(3)])
    e, comps = montar_registro(id_execucao="x", imovel_alvo={}, leilao_entrada={}, leilao={},
                               imoveis_coletados=[], resultado_ag2={}, zona_resultado={}, resultado_ag3={},
                               resultado_ag4={}, resultado_ag5=r)
    assert e["qtd_terrenos"] == 0 and e["valor_m2_terreno"] is None
    assert sum(1 for c in comps if c["tipo"] == "terreno" and not c["entrou_no_calculo"]) == 3
    resumo = {l[0]: l[1] for l in load_workbook(io.BytesIO(excel_resultado(e, comps)))["Resumo"].iter_rows(values_only=True)}
    assert resumo["Terrenos no cálculo"] == 0
    assert resumo["Terreno não separado"].startswith("Só 3 terreno(s) à venda na região")


def test_casa_sem_area_de_terreno_nao_usa_os_terrenos_da_zona():
    # Palhoca (20/10/2026): sem area do terreno o calculo e pelo m2 das casas, mas os 11
    # terrenos da zona apareciam como "usados" (funil, banco e Excel). O valor nao muda.
    from services.registro_execucao import montar_registro
    alvo = {"propertyType": "casa", "area": 150}
    terrenos = [terreno(200000 + i * 10000, 250 + i * 5) for i in range(12)]
    r = calcular(alvo, CASAS_5, terrenos)
    assert valor_medio(r) == valor_medio(calcular(alvo, CASAS_5, [])) == 700000.0
    assert r["comparaveis_usados"]["terreno"] == []
    descartados = r["comparaveis_descartados"]["terreno"]
    assert len(descartados) == 12 and {d["motivo_descarte"] for d in descartados} == {
        "terreno do imóvel não informado: o valor foi calculado pelo m² das casas"}
    e, linhas = montar_registro(id_execucao="x", imovel_alvo={}, leilao_entrada={}, leilao={},
                                imoveis_coletados=[], resultado_ag2={}, zona_resultado={}, resultado_ag3={},
                                resultado_ag4={}, resultado_ag5=r)
    assert e["qtd_terrenos"] == 0
    assert [l["entrou_no_calculo"] for l in linhas if l["tipo"] == "terreno"] == [False] * 12


def test_apartamento_nao_usa_terrenos():
    r = calcular({"propertyType": "apartamento", "area": 80},
                 [apto(400000 + i * 10000, 80) for i in range(5)], [terreno(200000, 250)])
    assert r["comparaveis_usados"]["terreno"] == []
    assert r["comparaveis_descartados"]["terreno"][0]["motivo_descarte"].startswith("apartamento/sala")


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
