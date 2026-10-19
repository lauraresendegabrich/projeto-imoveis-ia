"""
Testes do aviso de provavel lancamento (Agente 1 + alerta). A regra de repetidos
(agents/identidade.py) nao muda. Sem rede.

    .venv/Scripts/python.exe -m tests.test_lancamento
"""
import sys

sys.path.insert(0, ".")

from agents.collector import _marcar_provaveis_lancamentos
from agents.identidade import deduplicar
from services.alertas import calcular_alertas
from services.leilao import texto_aviso_lancamento

PORTAIS = ["vivareal", "imovelweb", "chavesnamao", "lugarcerto"]


def apto(i, preco=650000.0, area=72.0, anunciante=None, tipo="apartamento"):
    return {"portal": PORTAIS[i % 4], "source": PORTAIS[i % 4], "tipo": tipo, "preco": preco,
            "area_construida": area, "quartos": 2, "vagas": 1, "suites": 1, "bairro": "Centro",
            "rua": "Rua do Lancamento, 100", "anunciante_nome": anunciante or f"Imobiliaria {i}",
            "listing_id": str(i), "url": f"https://x/{i}", "descricao": ""}


def _agrupar(anuncios):
    unicos, descartados = deduplicar(anuncios)
    _marcar_provaveis_lancamentos(unicos, descartados)
    return unicos


def test_grupo_de_5_imobiliarias_diferentes_tem_alerta():
    unicos = _agrupar([apto(i) for i in range(6)])
    assert len(unicos) == 1                                   # continua 1 imovel
    marca = unicos[0]["provavel_lancamento"]
    assert marca == {"anuncios": 6, "imobiliarias": 6, "preco": 650000, "area": 72}, marca


def test_grupo_de_5_ou_mais_da_mesma_imobiliaria_sem_alerta():
    # Mesma imobiliaria em varios portais: e repeticao do mesmo anuncio, nao lancamento.
    unicos = _agrupar([apto(i, anunciante="Imobiliaria Unica") for i in range(6)])
    assert len(unicos) == 1 and "provavel_lancamento" not in unicos[0], unicos[0].get("provavel_lancamento")


def test_grupo_de_2_sem_alerta():
    unicos = _agrupar([apto(0), apto(1)])
    assert len(unicos) == 1 and "provavel_lancamento" not in unicos[0]


def test_grupo_com_menos_de_5_imobiliarias_sem_alerta():
    unicos = _agrupar([apto(i, anunciante="Imobiliaria A" if i < 3 else f"Imobiliaria {i}") for i in range(6)])
    assert all("provavel_lancamento" not in u for u in unicos)


def test_casa_nao_e_lancamento():
    unicos = _agrupar([apto(i, tipo="casa") for i in range(6)])
    assert all("provavel_lancamento" not in u for u in unicos)


def test_alerta_aviso_e_excel():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado, montar_registro
    marca = {"anuncios": 6, "imobiliarias": 6, "preco": 650000, "area": 72}
    ag5 = {"comparaveis_usados": {"construcao": [{"preco": 650000, "area": 72, "rua": "Rua do Lancamento, 100",
                                                  "status_zona": "confirmado", "provavel_lancamento": marca}],
                                  "terreno": []}}
    esperado = ("6 anúncios iguais (mesmo preço e área, imobiliárias diferentes; provável lançamento) "
                "foram contados como 1 imóvel.")
    assert texto_aviso_lancamento(ag5) == esperado
    alertas = calcular_alertas({"status_agente5": "ok", "qtd_comparaveis_construcao": 10}, ag5)
    lanc = [a for a in alertas if a["codigo"] == "provavel_lancamento"]
    assert len(lanc) == 1 and lanc[0]["mensagem"] == esperado, alertas
    e, comps = montar_registro(id_execucao="x", imovel_alvo={}, leilao_entrada={}, leilao={},
                               imoveis_coletados=[], resultado_ag2={}, zona_resultado={}, resultado_ag3={},
                               resultado_ag4={}, resultado_ag5=ag5)
    livro = load_workbook(io.BytesIO(excel_resultado(e, comps)))
    resumo = {l[0]: l[1] for l in livro["Resumo"].iter_rows(values_only=True)}
    assert resumo["Provável lançamento"] == esperado
    assert any(l[0] == "Provável lançamento" for l in livro["Alertas"].iter_rows(values_only=True))
    assert texto_aviso_lancamento({"comparaveis_usados": {"construcao": [{"preco": 1}]}}) is None


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
