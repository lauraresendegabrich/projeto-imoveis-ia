"""
Testes da ampliacao do raio da zona (Agente 2) ate ter 5 comparaveis, sem chamar a IA
de visao de novo. Sem rede.

    .venv/Scripts/python.exe -m tests.test_raio_ampliado
"""
import sys

sys.path.insert(0, ".")

import agents.comparables as ag2
from services.leilao import texto_aviso_raio
from tests._pasta_temporaria import data_dir_temporario

ALVO = (-23.30, -51.16)
M_POR_GRAU = 111_320.0


def anuncio(i, metros_ao_norte, tipo="Casas", **extra):
    """Anuncio com coordenada do portal a `metros_ao_norte` do alvo."""
    return {"id": i, "url": f"https://x/{tipo}/{i}", "propertyType": tipo, "price": 300000 + i, "area": 100,
            "cluster": "A", "neighborhood": f"Bairro {i}", "lat": ALVO[0] + metros_ao_norte / M_POR_GRAU,
            "lon": ALVO[1], **extra}


def fora(i, dist, **extra):
    return {"id": i, "distancia_metros": dist, "geocodificacao_nivel": "coordenadas_athena", **extra}


# ---------------------------------------------------------------- _ampliar_raio
def test_sem_ampliacao_quando_ja_tem_5():
    conf = [fora(i, 100) for i in range(5)]
    c, f, raio = ag2._ampliar_raio(conf, [fora(9, 600)], 450)
    assert raio == 450 and len(c) == 5 and len(f) == 1


def test_ampliacao_ate_chegar_a_5():
    # 450 -> 675 (entram 500 e 600) -> 1012 (entra 900) = 5
    conf = [fora(1, 100), fora(2, 200)]
    c, f, raio = ag2._ampliar_raio(conf, [fora(3, 500), fora(4, 600), fora(5, 900), fora(6, 1400)], 450)
    assert raio == 1012 and len(c) == 5, (raio, len(c))
    assert {im["id"] for im in f} == {6}
    assert all(im.get("incluido_por_ampliacao_raio") for im in c[2:])


def test_ampliacao_bate_no_limite_sem_chegar_a_5():
    c, f, raio = ag2._ampliar_raio([fora(1, 100)], [fora(2, 1200), fora(3, 1600), fora(4, 3000)], 450)
    assert raio == 1500 and len(c) == 2 and {im["id"] for im in f} == {3, 4}, (raio, len(c))


def test_raio_da_ia_maior_que_o_limite_nao_amplia():
    c, f, raio = ag2._ampliar_raio([fora(1, 100)], [fora(2, 1800)], 2000)
    assert raio == 2000 and len(c) == 1


def test_nao_amplia_com_distancia_do_bairro_ou_coordenada_generica():
    lista = [fora(2, 500, geocodificacao_nivel="bairro"),
             fora(3, 500, motivo_zona_nao_verificada="coordenada generica (a 10 m do centro da cidade)"),
             {"id": 4, "distancia_metros": None}]
    c, f, raio = ag2._ampliar_raio([fora(1, 100)], lista, 450)
    assert len(c) == 1 and raio == 1500


# ---------------------------------------------------------------- zona completa
def _zona(casas, terrenos):
    chamadas = {"ia": 0, "geo": 0}

    def ia(*a, **k):
        chamadas["ia"] += 1
        return {"raio_metros": 450, "raio_sugerido_metros": 450}

    def geo(**k):
        chamadas["geo"] += 1
        return None, None, None, None

    originais = (ag2._obter_imagem_satelite, ag2._analisar_zona_homogenea, ag2._centros_da_cidade,
                 ag2._geocodificar_com_fallback)
    ag2._obter_imagem_satelite = lambda *a, **k: b"png"
    ag2._analisar_zona_homogenea = ia
    ag2._centros_da_cidade = lambda *a, **k: []
    ag2._geocodificar_com_fallback = geo
    try:
        with data_dir_temporario():
            z = ag2.analisar_zona_homogenea("x", casas, cidade="Londrina", estado="PR", lat_alvo_precomp=ALVO[0],
                                            lon_alvo_precomp=ALVO[1], terrenos=terrenos, id_execucao="teste")
    finally:
        (ag2._obter_imagem_satelite, ag2._analisar_zona_homogenea, ag2._centros_da_cidade,
         ag2._geocodificar_com_fallback) = originais
    return z, chamadas


def test_zona_amplia_casas_e_terrenos_separadamente_e_chama_a_ia_uma_vez():
    casas = [anuncio(i, 100 + 50 * i) for i in range(3)] + [anuncio(10, 600), anuncio(11, 1300)]
    terrenos = [anuncio(20 + i, 200 + 40 * i, "Terrenos") for i in range(5)] + [anuncio(30, 1400, "Terrenos")]
    z, chamadas = _zona(casas, terrenos)
    assert chamadas == {"ia": 1, "geo": 0}, chamadas
    # 450 -> 675 (entra a de 600 m) -> 1012 -> 1500 (entra a de 1300 m) = 5
    assert z["raio_ia_m"] == 450 and z["raio_final_m"] == 1500, z["raio_final_m"]
    assert z["raio_final_terrenos_m"] == 450                        # terrenos ja tinham 5
    assert z["resumo_zona"]["raio_usado_metros"] == z["raio_final_m"]
    assert z["resumo_zona"]["confirmados"] == 5 and z["ampliacao_raio"]["confirmados_antes"] == 3
    aviso = texto_aviso_raio(z)
    assert aviso == ("Havia menos de 5 comparáveis no raio de 450 m; o raio foi ampliado para "
                     f"{z['raio_final_m']:,} m para chegar a 5.".replace(",", ".")), aviso


def test_zona_no_limite_sem_completar_5():
    casas = [anuncio(i, 100 * (i + 1)) for i in range(2)] + [anuncio(10, 2500)]
    z, _ = _zona(casas, [])
    assert z["raio_final_m"] == 1500 and z["resumo_zona"]["confirmados"] == 2
    assert texto_aviso_raio(z) == ("Havia menos de 5 comparáveis no raio de 450 m; o raio foi ampliado até "
                                   "1.500 m e ainda assim só 2 foram confirmados."), texto_aviso_raio(z)


def test_zona_sem_ampliacao_sem_aviso():
    z, chamadas = _zona([anuncio(i, 50 * (i + 1)) for i in range(5)], [])
    assert z["raio_final_m"] == 450 and not z["ampliacao_raio"]["ampliado"] and texto_aviso_raio(z) is None
    assert chamadas["ia"] == 1


# ---------------------------------------------------------------- banco e Excel
def test_registro_e_excel_com_os_dois_raios():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado, montar_registro
    casas = [anuncio(i, 100 * (i + 1)) for i in range(2)] + [anuncio(10, 2500)]
    z, _ = _zona(casas, [])
    e, comps = montar_registro(id_execucao="x", imovel_alvo={}, leilao_entrada={}, leilao={},
                               imoveis_coletados=[], resultado_ag2={}, zona_resultado=z, resultado_ag3={},
                               resultado_ag4={}, resultado_ag5={})
    assert e["raio_zona_m"] == 1500
    assert e["parametros"]["raio_ia_m"] == 450 and e["parametros"]["raio_final_m"] == 1500
    resumo = {l[0]: l[1] for l in load_workbook(io.BytesIO(excel_resultado(e, comps)))["Resumo"].iter_rows(values_only=True)}
    assert resumo["Raio escolhido pela IA (m)"] == 450 and resumo["Raio usado no cálculo (m)"] == 1500
    assert resumo["Ampliação do raio"].startswith("Havia menos de 5 comparáveis no raio de 450 m")


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
