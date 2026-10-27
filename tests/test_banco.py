"""
Testes do registro no banco (services/registro_execucao.py e services/banco.py).
Sem banco real: confere a montagem das linhas, a coerencia com db/schema.sql e que
uma falha do banco nunca quebra nem trava.

    .venv/Scripts/python.exe -m tests.test_banco
"""
import re
import sys
import time

sys.path.insert(0, ".")

from services.banco import (
    COLUNAS_COMPARAVEIS, COLUNAS_EXECUCOES, COLUNAS_IMOVEIS_LEILAO, gravar_execucao, sql_gravar_imovel_leilao,
)
from services.leilao import avaliar_leilao
from services.registro_execucao import montar_registro, versao_codigo


def _ag5():
    return {
        "status_avaliacao": "ok",
        "avaliacao_confiavel": True,
        "zona_calculada_nesta_avaliacao": True,
        "avaliacao_planilha": {"valor_medio_imovel": 716036.0, "valor_liquidez": 644432.4},
        "valor_m2_zona_homogenea": {
            "terreno": {"valor_m2_referencia": 900.0},
            "construcao": {"combinados": {"valor_m2_referencia": 3500.0}},
        },
        "liquidez_experimental": {"tempo_estimado": "60 a 90 dias"},
        "comparaveis_usados": {
            "construcao": [
                {"url": "https://a/1", "rua": "Rua A, 10", "bairro": "Centro", "preco": 500000, "area": 100,
                 "status_zona": "confirmado", "portal": "VivaReal"},
                {"url": "https://a/2", "rua": "Rua B", "bairro": "Centro", "preco": 400000, "area": 80,
                 "status_zona": "fallback", "portal": "ImovelWeb"},
            ],
            "terreno": [{"url": "https://t/1", "preco": 200000, "area": 250, "status_zona": "confirmado"}],
        },
        "comparaveis_descartados": {
            "construcao": [{"url": "https://a/3", "preco": 9000000, "area": 100, "status_zona": "confirmado"}],
            "terreno": [],
        },
        "avisos": [],
    }


def _registro(**extra):
    ag5 = _ag5()
    entrada = {"origem": "CAIXA", "valor_minimo": 323158, "valor_avaliacao": 410000,
               "modalidade": "Venda Online", "codigo_caixa": "1444400000000", "link": "https://caixa/1"}
    alvo = {"tipo": "house", "rua": "Rua X", "numero": "100", "bairro": "Centro", "cidade": "Curitiba",
            "estado": "PR", "area": 150, "area_terreno": 300, "bedrooms": 3, "bathrooms": 2}
    coletados = [{"fonte_dados": "Athena/S3", "duplicatas": [{}, {}]}, {"fonte_dados": "Athena/S3"}, {}]
    args = dict(
        id_execucao="abc123", imovel_alvo=alvo, leilao_entrada=entrada,
        leilao=avaliar_leilao(ag5, 323158, 410000), imoveis_coletados=coletados,
        resultado_ag2={"resumo": {"cluster_a": 12}, "comparaveis": [
            {"status_julgamento": "JULGADO_LLM", "provider_llm": "groq", "modelo_llm": "gpt-oss-120b"}]},
        zona_resultado={"resumo_zona": {"raio_usado_metros": 450, "confirmados": 9},
                        "zona_homogenea": {"provider_visao": "gemini"}},
        resultado_ag3={}, resultado_ag4={"scores": {"provider_llm": "Gemini"}},
        resultado_ag5=ag5, tempos={"total": 300.5}, falhas=["agente3: Timeout"],
    )
    args.update(extra)
    return montar_registro(**args)


def test_execucao_tem_os_campos_principais():
    e, _ = _registro()
    assert e["origem"] == "CAIXA" and e["tipo"] == "casa" and e["uf"] == "PR"
    assert e["endereco"] == "Rua X, 100"
    assert e["valor_minimo"] == 323158 and e["valor_avaliacao"] == 410000
    assert e["decisao"] == "Oportunidade" and round(e["lance_maximo"]) == 337068 and round(e["sobra"]) == 13910
    assert e["status_agente5"] == "ok" and e["liquidez"] == 644432.4 and e["valor_mercado"] == 716036.0
    assert e["anuncios_encontrados"] == 3 and e["repetidos"] == 2 and e["parecidos"] == 12
    assert e["confirmados_na_zona"] == 9 and e["raio_zona_m"] == 450
    assert e["usados_por_fallback"] == 1 and e["terrenos_usados"] == 1
    assert e["qtd_comparaveis_construcao"] == 2 and e["qtd_terrenos"] == 1
    assert e["tempo_venda"] == "60 a 90 dias" and e["tempo_s"] == 300.5
    assert e["parametros"]["regra_leilao"]["roi_oportunidade"] == 0.30 and e["parametros"]["min_amostras_confiavel"] == 5
    assert e["modelos_ia"]["agente4"] == "Gemini" and e["fontes"]["anuncios_apify"] == 1
    assert e["falhas"]["falhas"] == ["agente3: Timeout"]


def test_sem_leilao_decisao_fica_vazia():
    e, _ = _registro(leilao=None, leilao_entrada=None)
    assert e["decisao"] is None and e["lance_maximo"] is None and e["origem"] == "CAIXA"


def test_comparaveis_usados_e_descartados():
    _, linhas = _registro()
    assert len(linhas) == 4, linhas
    assert sum(1 for l in linhas if l["entrou_no_calculo"]) == 3
    descartado = next(l for l in linhas if not l["entrou_no_calculo"])
    assert descartado["link"] == "https://a/3" and descartado["valor_m2"] == 90000.0
    fallback = next(l for l in linhas if l["link"] == "https://a/2")
    assert fallback["status_zona"] == "fallback" and fallback["portal"] == "ImovelWeb" and fallback["valor_m2"] == 5000.0
    assert {l["tipo"] for l in linhas} == {"construcao", "terreno"}
    assert all(l["id_execucao"] == "abc123" for l in linhas)


def test_linhas_tem_exatamente_as_colunas_gravadas():
    e, linhas = _registro()
    assert set(COLUNAS_EXECUCOES) <= set(e), set(COLUNAS_EXECUCOES) - set(e)
    assert all(set(COLUNAS_COMPARAVEIS) <= set(l) for l in linhas)


def test_colunas_batem_com_schema_sql():
    sql = open("db/schema.sql", encoding="utf-8").read()

    def colunas(tabela):
        corpo = re.search(rf"CREATE TABLE IF NOT EXISTS {tabela} \((.*?)\n\);", sql, re.S).group(1)
        nomes = set()
        for linha in corpo.splitlines():
            linha = linha.split("--")[0].strip()
            m = re.match(r"([a-z_0-9]+)\s+[A-Z]", linha)
            if m:
                nomes.add(m.group(1))
        return nomes

    no_schema_exec = colunas("execucoes") - {"data_hora"}
    assert no_schema_exec == set(COLUNAS_EXECUCOES), (no_schema_exec ^ set(COLUNAS_EXECUCOES))
    no_schema_comp = colunas("comparaveis") - {"id"}
    assert no_schema_comp == set(COLUNAS_COMPARAVEIS), (no_schema_comp ^ set(COLUNAS_COMPARAVEIS))


def test_sem_database_url_nao_grava_e_nao_quebra():
    e, linhas = _registro()
    assert gravar_execucao(e, linhas, database_url="") is False


def test_banco_fora_do_ar_nao_quebra_nem_trava():
    e, linhas = _registro()
    t0 = time.time()
    ok = gravar_execucao(e, linhas, database_url="postgresql://u:p@127.0.0.1:1/db?connect_timeout=3")
    assert ok is False and time.time() - t0 < 20


def _abrir_excel():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado
    e, comps = _registro()
    return load_workbook(io.BytesIO(excel_resultado(e, comps))), comps


def test_excel_tem_abas_resumo_e_comparaveis():
    livro, _ = _abrir_excel()
    assert livro.sheetnames == ["Resumo", "Alertas", "Comparáveis"], livro.sheetnames


def test_excel_resumo_campo_valor():
    livro, _ = _abrir_excel()
    resumo = {linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True)}
    assert resumo["Decisão"] == "Oportunidade" and round(resumo["ROI pagando o valor mínimo"], 4) == 0.3495
    assert round(resumo["Lance máximo (ROI de 30%)"]) == 337068
    assert round(resumo["Sobra até o lance máximo (negativo = falta)"]) == 13910
    assert resumo["Valor mínimo CAIXA"] == 323158 and resumo["Valor de liquidez (−10%)"] == 644432.4
    assert resumo["Cidade"] == "Curitiba" and resumo["Estimativa confiável"] == "Sim"
    assert resumo["Identificador da avaliação"] == "abc123"
    celula_lance = next(l[1] for l in livro["Resumo"].iter_rows() if l[0].value == "Lance máximo (ROI de 30%)")
    assert "R$" in celula_lance.number_format


def test_excel_alertas_com_motivo_e_anuncios():
    livro, _ = _abrir_excel()
    resumo = {linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True)}
    linhas = list(livro["Alertas"].iter_rows(values_only=True))
    assert linhas[0][:2] == ("Alerta", "Motivo")
    nomes = [l[0] for l in linhas[1:] if l[0]]
    assert "Distância não verificada" in nomes and "Falha de agente" in nomes, nomes
    assert resumo["Alertas"] == f"{len(nomes)} (veja a aba Alertas)"
    # o anuncio fora da zona aparece logo abaixo do alerta, com link clicavel
    i = next(i for i, l in enumerate(linhas) if l[0] == "Distância não verificada")
    assert linhas[i + 1][1].startswith("↳ Rua B") and linhas[i + 1][5] == "https://a/2"
    assert livro["Alertas"].cell(row=i + 2, column=6).hyperlink is not None
    # o erro do agente aparece abaixo da falha
    j = next(j for j, l in enumerate(linhas) if l[0] == "Falha de agente")
    assert linhas[j + 1][1] == "↳ agente3: Timeout"


def test_excel_sem_alertas():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado
    e, comps = _registro(falhas=[])
    e["alertas"] = []
    livro = load_workbook(io.BytesIO(excel_resultado(e, comps)))
    resumo = {linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True)}
    assert resumo["Alertas"] == "Nenhum"
    assert list(livro["Alertas"].iter_rows(values_only=True))[1][0] == "Nenhum alerta nesta avaliação."


def test_excel_comparaveis_um_por_linha_com_link():
    livro, comps = _abrir_excel()
    aba = livro["Comparáveis"]
    linhas = list(aba.iter_rows(values_only=True))
    assert linhas[0][:3] == ("Tipo", "Entrou no cálculo", "Localização")
    corpo = linhas[1:]
    assert len(corpo) == len(comps)
    assert [l[0] for l in corpo] == ["Construção", "Construção", "Construção", "Terreno"]
    assert [l[1] for l in corpo[:3]] == ["Sim", "Sim", "Não"]   # usados antes dos descartados
    assert "Distância não verificada" in [l[2] for l in corpo]
    assert aba.cell(row=2, column=10).hyperlink is not None
    assert aba.freeze_panes == "A2"


def _registro_com_localizacao_e_leilao():
    from datetime import date
    entrada = {"origem": "CAIXA", "valor_minimo": 323158, "data_leilao": date(2026, 11, 25),
               "formas_pagamento": ["Recursos próprios (à vista)", "FGTS"],
               "regras_despesas": "Condomínio: sob responsabilidade do comprador."}
    coletados = [{"escopo_busca": "rua+bairro+proximidade"},
                 {"escopo_busca": "rua+bairro+proximidade", "origem_busca": "proximidade"},
                 {"escopo_busca": "rua+bairro+proximidade", "origem_busca": "proximidade"}]
    return _registro(leilao_entrada=entrada, imoveis_coletados=coletados, localizacao_alvo="rua_numero")


def test_grava_localizacao_busca_e_dados_do_leilao():
    from datetime import date
    e, _ = _registro_com_localizacao_e_leilao()
    assert e["localizacao_alvo"] == "rua_numero"
    assert e["escopo_busca"] == "rua+bairro+proximidade" and e["anuncios_por_proximidade"] == 2
    assert e["data_leilao"] == date(2026, 11, 25)
    assert e["formas_pagamento"] == ["Recursos próprios (à vista)", "FGTS"]
    assert e["regras_despesas"].startswith("Condomínio")


def test_sem_dados_novos_ficam_vazios():
    e, _ = _registro(leilao_entrada=None, imoveis_coletados=[])
    assert e["localizacao_alvo"] is None and e["escopo_busca"] is None
    assert e["anuncios_por_proximidade"] == 0
    assert e["data_leilao"] is None and e["formas_pagamento"] is None and e["regras_despesas"] is None


def test_excel_mostra_localizacao_e_dados_do_leilao():
    import io
    from datetime import datetime
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado
    e, comps = _registro_com_localizacao_e_leilao()
    livro = load_workbook(io.BytesIO(excel_resultado(e, comps)))
    resumo = {linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True)}
    assert resumo["Localização do imóvel no mapa"] == "Rua e número"
    assert resumo["Onde os anúncios foram buscados"] == "Rua, bairro e proximidade"
    assert resumo["Anúncios trazidos por proximidade (até 1,5 km)"] == 2
    assert resumo["Data do leilão"] == datetime(2026, 11, 25)
    assert resumo["Formas de pagamento"] == "Recursos próprios (à vista), FGTS"
    assert resumo["Regras de despesas"].startswith("Condomínio")
    e["localizacao_alvo"] = "nao_localizado"
    livro = load_workbook(io.BytesIO(excel_resultado(e, comps)))
    resumo = {linha[0]: linha[1] for linha in livro["Resumo"].iter_rows(values_only=True)}
    assert resumo["Localização do imóvel no mapa"].startswith("Não localizado")


def test_migracao_002_cria_as_colunas_novas():
    sql = open("db/migracoes/002_localizacao_e_dados_do_leilao.sql", encoding="utf-8").read()
    for coluna in ("localizacao_alvo", "escopo_busca", "anuncios_por_proximidade",
                   "data_leilao", "formas_pagamento", "regras_despesas"):
        assert f"ADD COLUMN IF NOT EXISTS {coluna}" in sql, coluna
        assert coluna in COLUNAS_EXECUCOES, coluna


# ---------------------------------------------------------------- imoveis_leilao (migracao 005)
def _colunas_da_tabela(sql, tabela):
    corpo = re.search(rf"CREATE TABLE IF NOT EXISTS {tabela} \((.*?)\n\);", sql, re.S).group(1)
    nomes = set()
    for linha in corpo.splitlines():
        m = re.match(r"([a-z_0-9]+)\s+[A-Z]", linha.split("--")[0].strip())
        if m:
            nomes.add(m.group(1))
    return nomes


def test_imoveis_leilao_bate_com_schema_e_migracao():
    automaticas = {"id", "situacao", "criado_em", "atualizado_em"}
    for arquivo in ("db/schema.sql", "db/migracoes/005_imoveis_leilao.sql"):
        sql = open(arquivo, encoding="utf-8").read()
        assert _colunas_da_tabela(sql, "imoveis_leilao") - automaticas == set(COLUNAS_IMOVEIS_LEILAO), arquivo
        assert "UNIQUE (banco, codigo)" in sql and "CREATE OR REPLACE VIEW vw_leiloes_ativos" in sql, arquivo
    migracao = open("db/migracoes/005_imoveis_leilao.sql", encoding="utf-8").read()
    for coluna in ("id_imovel_leilao", "etapa_avaliada", "roi", "lance_limite_descarte"):
        assert f"ADD COLUMN IF NOT EXISTS {coluna}" in migracao and coluna in COLUNAS_EXECUCOES, coluna


def test_codigo_do_imovel_igual_no_numero_e_no_link():
    from services.registro_execucao import codigo_imovel
    assert codigo_imovel("855552271653-1") == codigo_imovel(None, "https://x/detalhe-imovel.asp?hdnimovel=8555522716531")
    assert codigo_imovel("000001015142-0") == "10151420" == codigo_imovel("", "https://x?hdnimovel=10151420")
    assert codigo_imovel(None) is None and codigo_imovel("", "https://x/sem-codigo") is None


def test_execucao_traz_roi_e_o_imovel_de_leilao():
    e, _ = _registro()
    assert round(e["roi"], 4) == 0.3495 and round(e["lance_limite_descarte"]) == 387629
    assert e["etapa_avaliada"] is None and e["id_imovel_leilao"] is None   # o id vem na gravacao
    imovel = e["_imovel_leilao"]
    assert set(imovel) == set(COLUNAS_IMOVEIS_LEILAO)
    assert imovel["codigo"] == "1444400000000" and imovel["modalidade"] == "Venda Online"
    assert imovel["valor_minimo"] == 323158 and imovel["cidade"] == "Curitiba" and imovel["uf"] == "PR"
    assert imovel["valor_minimo_1"] is None and imovel["gravame"] is None   # nada inventado


def test_data_sem_hora_vale_ate_o_fim_do_dia():
    from datetime import date, datetime, time
    entrada = {"modalidade": "Licitação Aberta", "codigo_caixa": "178770097494-1", "data_leilao": date(2026, 11, 4)}
    e, _ = _registro(leilao_entrada=entrada)
    assert e["_imovel_leilao"]["data_encerramento"] == datetime.combine(date(2026, 11, 4), time.max)


def test_leilao_sfi_com_as_duas_etapas():
    from datetime import datetime
    entrada = {"modalidade": "Leilão SFI", "numero_imovel": "855552271653-1",
               "valor_minimo_1": 224000.00, "data_1_leilao": datetime(2026, 11, 30, 10, 0),
               "valor_minimo_2": 296523.50, "data_2_leilao": datetime(2026, 12, 4, 10, 0),
               "edital": "0052/0226 - CPA/RE", "numero_item": "25", "gravame": False}
    imovel = _registro(leilao_entrada=entrada)[0]["_imovel_leilao"]
    assert imovel["codigo"] == "8555522716531" and imovel["numero_imovel"] == "855552271653-1"
    assert (imovel["valor_minimo_1"], imovel["valor_minimo_2"]) == (224000.00, 296523.50)
    assert imovel["data_2_leilao"] == datetime(2026, 12, 4, 10, 0) and imovel["data_encerramento"] is None
    assert imovel["edital"] == "0052/0226 - CPA/RE" and imovel["gravame"] is False



def test_execucao_grava_o_valor_e_a_data_da_etapa_avaliada():
    from datetime import date, datetime
    from services.leilao import avaliar_imovel_leilao
    entrada = {"modalidade": "Leilão SFI", "numero_imovel": "855552271653-1",
               "valor_minimo_1": 224000.00, "data_1_leilao": datetime(2026, 11, 30, 10, 0),
               "valor_minimo_2": 296523.50, "data_2_leilao": datetime(2026, 12, 4, 10, 0)}
    leilao = avaliar_imovel_leilao(_ag5(), entrada, agora=datetime(2026, 12, 1, 9, 0))
    e, _ = _registro(leilao_entrada=entrada, leilao=leilao)
    assert e["etapa_avaliada"] == 2 and e["valor_minimo"] == 296523.50
    assert e["data_leilao"] == date(2026, 12, 4) and e["decisao"] == leilao["decisao"]

def test_sem_codigo_ou_modalidade_antiga_nao_grava_o_imovel():
    for entrada in ({"modalidade": "Venda Online"}, {"modalidade": "Outra", "codigo_caixa": "123"},
                    {"modalidade": None, "codigo_caixa": "123"}, None):
        assert _registro(leilao_entrada=entrada)[0]["_imovel_leilao"] is None, entrada


def test_sql_do_imovel_nao_apaga_dados_nem_a_situacao():
    sql = sql_gravar_imovel_leilao()
    assert "ON CONFLICT (banco, codigo) DO UPDATE" in sql and "RETURNING id" in sql
    assert "modalidade = COALESCE(EXCLUDED.modalidade, imoveis_leilao.modalidade)" in sql
    assert "situacao" not in sql and "codigo = " not in sql


class _CursorFalso:
    def __init__(self, log):
        self.log = log

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def execute(self, sql, parametros=None):
        self.log.append((sql.split()[0] + " " + sql.split()[2], parametros))

    def executemany(self, sql, linhas):
        self.log.append(("INSERT comparaveis", linhas))

    def fetchone(self):
        return (42,)


class _ConexaoFalsa:
    def __init__(self, log):
        self.log = log

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def transaction(self):
        return self

    def cursor(self):
        return _CursorFalso(self.log)


def test_grava_o_imovel_antes_e_liga_a_execucao_pelo_id():
    import psycopg
    log = []
    original = psycopg.connect
    psycopg.connect = lambda *a, **k: _ConexaoFalsa(log)
    try:
        e, linhas = _registro()
        assert gravar_execucao(e, linhas, database_url="postgresql://falso") is True
    finally:
        psycopg.connect = original
    comandos = [c for c, _ in log if not c.startswith("SET")]
    assert comandos == ["INSERT imoveis_leilao", "INSERT execucoes", "INSERT comparaveis"], comandos
    parametros_execucao = next(p for c, p in log if c == "INSERT execucoes")
    assert parametros_execucao[COLUNAS_EXECUCOES.index("id_imovel_leilao")] == 42
    assert e["id_imovel_leilao"] is None   # o dict original nao e alterado


def test_sem_imovel_de_leilao_grava_so_a_execucao():
    import psycopg
    log = []
    original = psycopg.connect
    psycopg.connect = lambda *a, **k: _ConexaoFalsa(log)
    try:
        e, linhas = _registro(leilao_entrada={"modalidade": "Venda Online"})
        assert gravar_execucao(e, linhas, database_url="postgresql://falso") is True
    finally:
        psycopg.connect = original
    assert [c for c, _ in log if not c.startswith("SET")] == ["INSERT execucoes", "INSERT comparaveis"]


def test_versao_codigo_do_git():
    v = versao_codigo()
    assert v == "desconhecida" or re.fullmatch(r"[0-9a-f]{40}", v), v


if __name__ == "__main__":
    import logging
    logging.disable(logging.CRITICAL)
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
