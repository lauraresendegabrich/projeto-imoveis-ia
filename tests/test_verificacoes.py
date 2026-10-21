"""
Verificacoes automaticas (services/verificacoes.py): para cada uma, um caso que
passa e um que falha. O resultado de base usa o Agente 5 de verdade. Sem rede.

    .venv/Scripts/python.exe -m tests.test_verificacoes
"""
import copy
import logging
import sys

sys.path.insert(0, ".")
logging.disable(logging.CRITICAL)

from agents.price_liquidity import executar_agente5
from services.verificacoes import ordenar_para_tela, refazer_conta_agente5, verificar_execucao

ID = "exec-1"


def casa(i, **extra):
    return {"url": f"https://x/{i}", "propertyType": "Casas", "price": 500_000 + i * 37_000, "area": 100 + i * 4,
            "street": f"Rua {i}", "neighborhood": "Centro", "city": "Curitiba", "state": "PR", "cluster": "A",
            "status_julgamento": "JULGADO_LLM", "provider_llm": "gemini", "modelo_llm": "gemini-2.5-flash",
            "checagem_alvo": {"resultado": "nao", "score": 0.0, "sinais": []},
            "classificacao_zona": "na_zona", "distancia_metros": 100 + i * 50,
            "id_execucao": ID, "coletado_em": "2026-10-21T09:00:00", **extra}


def terreno(i, **extra):
    return {"url": f"https://t/{i}", "propertyType": "Terrenos", "price": 200_000 + i * 15_000, "area": 300 + i * 10,
            "city": "Curitiba", "state": "PR", "cluster": "terreno", "status_julgamento": "SEPARADO_TIPO",
            "checagem_alvo": {"resultado": "nao", "score": 0.0, "sinais": []},
            "classificacao_zona": "na_zona", "distancia_metros": 200 + i * 30,
            "id_execucao": ID, "coletado_em": "2026-10-21T09:00:00", **extra}


def montar(area_terreno=None, n_casas=6, n_terrenos=2, extras_terreno=None):
    casas = [casa(i) for i in range(n_casas)]
    reprovada = casa(90, cluster="B", status_julgamento="REPROVADO_PRE_CLASSIFICACAO",
                     classificacao_zona=None, distancia_metros=None)
    terrenos = [terreno(i, **(extras_terreno or {}).get(i, {})) for i in range(n_terrenos)]
    alvo5 = {"propertyType": "casa", "area": 120, "area_terreno": area_terreno}
    ag5 = executar_agente5(alvo5, copy.deepcopy(terrenos), copy.deepcopy(casas), {}, {})
    ag5["id_execucao"] = ID
    ag5["zona_calculada_nesta_avaliacao"] = True
    return {
        "id_execucao": ID,
        "imovel_alvo": {"cidade": "Curitiba", "estado": "PR", "tipo": "house", "area": 120,
                        "geocodificacao_nivel": "rua"},
        "localizacao_alvo": "rua",
        "resumo_coleta": {"fontes": {"athena": {"consultada": True, "respondeu": True, "anuncios": 9, "falhas": []},
                                     "apify": {"consultada": False, "motivo": "não consultado"}},
                          "contagem": {"antes_dos_repetidos": n_casas + n_terrenos + 2, "unicos": n_casas + n_terrenos + 1,
                                       "repetidos_juntados": 1, "cortados_pelo_limite": 0, "repetidos_dos_cortados": 0,
                                       "mantidos": n_casas + n_terrenos + 1, "repetidos": 1}},
        "falhas": [],
        "comparaveis": casas + [reprovada],
        "terrenos": terrenos,
        "resumo": {"cluster_a": n_casas, "llm_nao_julgados": 0},
        "duplicatas_descartadas": [{"Mantido": "VivaReal", "Preço": 500000, "Área": 100, "Descartado": "ImovelWeb",
                                    "Anúncio descartado": "https://iw/1", "Motivo": "mesmo_codigo"}],
        "zona_homogenea": {"id_execucao": ID, "coordenadas_alvo": {"lat": -25.4, "lon": -49.2},
                           "raio_final_m": 500, "comparaveis_confirmados": casas + terrenos, "terrenos": terrenos},
        "analise_qualitativa": {"id_execucao": ID,
                                "imovel_alvo": {"analise_qualitativa": {"status": "ok", "scores": {"score_qualitativo": 0.5}}},
                                "comparaveis": [dict(c, analise_qualitativa={"status": "ok", "scores": {"score_qualitativo": 0.6}})
                                                for c in casas]},
        "infraestrutura": {"id_execucao": ID, "scores": {"score_final": 0.72}},
        "preco_estimado": ag5,
    }


def _v(resultado, inicio):
    """A verificacao cujo texto comeca com `inicio`."""
    achadas = [v for v in verificar_execucao(resultado) if v["verificacao"].startswith(inicio)]
    assert len(achadas) == 1, (inicio, [v["verificacao"] for v in verificar_execucao(resultado)])
    return achadas[0]


def _passa(resultado, inicio):
    v = _v(resultado, inicio)
    assert v["ok"] is True, v


def _falha(resultado, inicio):
    v = _v(resultado, inicio)
    assert v["ok"] is False, v
    return v


# ---------------------------------------------------------------- base
def test_base_tudo_passa():
    vs = verificar_execucao(montar())
    assert all(v["ok"] is True for v in vs), [v for v in vs if v["ok"] is not True]
    assert all(set(v) == {"agente", "verificacao", "esperado", "encontrado", "ok", "problema"} for v in vs)
    assert all(v["problema"] is None for v in vs)


def test_falhas_primeiro_na_tela():
    r = montar()
    r["preco_estimado"]["id_execucao"] = "outra"
    r.pop("resumo_coleta")                     # coleta antiga: fontes e contagens nao verificaveis
    ordem = [v["ok"] for v in ordenar_para_tela(verificar_execucao(r))]
    assert ordem == sorted(ordem, key=lambda ok: {False: 0, "alerta": 1, None: 2, True: 3}[ok]) and ordem[0] is False
    assert None in ordem and ordem[-1] is True


def test_resultado_vazio_nao_quebra():
    assert isinstance(verificar_execucao({}), list)


# ---------------------------------------------------------------- geral
def test_geral_funil_uma_etapa_por_anuncio():
    _passa(montar(), "Cada anúncio aparece em exatamente uma etapa")
    r = montar()
    r["comparaveis"].append(dict(r["comparaveis"][0]))     # mesmo anuncio duas vezes
    v = _falha(r, "Cada anúncio aparece em exatamente uma etapa")
    assert "1 em mais de uma etapa" in v["encontrado"], v


def test_geral_resultados_desta_execucao():
    _passa(montar(), "Os resultados dos Agentes 2 a 5 são desta execução")
    r = montar()
    r["infraestrutura"]["id_execucao"] = "rodada-anterior"
    assert "Agente 4: rodada-anterior" in _falha(r, "Os resultados dos Agentes 2 a 5")["encontrado"]


# ---------------------------------------------------------------- agente 1
def test_ag1_cidade_uf():
    _passa(montar(), "Todos os anúncios são da cidade/UF pedidas")
    r = montar()
    r["comparaveis"][0]["city"] = "Colombo"
    assert "1 de outra cidade (Colombo/PR)" in _falha(r, "Todos os anúncios são da cidade")["encontrado"]


# ---------------------------------------------------------------- agente 2
def test_ag2_decisao_de_cada_imovel():
    v = _v(montar(), "Todo imóvel construído tem decisão")
    assert v["ok"] and "1 pela pré-classificação, 6 pela IA (gemini/gemini-2.5-flash: 6)" in v["encontrado"], v
    r = montar()
    r["comparaveis"][1].update(status_julgamento="NAO_JULGADO", cluster=None)
    assert "1 sem julgamento" in _falha(r, "Todo imóvel construído tem decisão")["encontrado"]


def test_ag2_distancia_dos_confirmados():
    _passa(montar(), "Todo comparável confirmado na zona tem distância")
    r = montar()
    r["zona_homogenea"]["comparaveis_confirmados"][2]["distancia_metros"] = 900
    assert "1 fora do raio" in _falha(r, "Todo comparável confirmado na zona")["encontrado"]


def test_ag2_checagem_do_proprio_anuncio():
    _passa(montar(), "A checagem do próprio anúncio rodou")
    r = montar()
    del r["terrenos"][0]["checagem_alvo"]
    assert "1 sem" in _falha(r, "A checagem do próprio anúncio rodou")["encontrado"]


# ---------------------------------------------------------------- agente 3
def test_ag3_avaliados_e_notas():
    _passa(montar(), "Todo comparável que chegou foi avaliado")
    r = montar()
    r["analise_qualitativa"]["comparaveis"][0]["analise_qualitativa"]["scores"]["score_qualitativo"] = 1.4
    _falha(r, "Todo comparável que chegou foi avaliado")
    r = montar()
    r["analise_qualitativa"]["comparaveis"].pop()          # um nao avaliado, sem falha registrada
    _falha(r, "Todo comparável que chegou foi avaliado")
    r["analise_qualitativa"] = {}
    r["falhas"] = ["agente3: TimeoutError"]                 # falha registrada: ok
    _passa(r, "Todo comparável que chegou foi avaliado")


# ---------------------------------------------------------------- agente 4
def test_ag4_localizacao_e_nota():
    _passa(montar(), "Rodou com localização precisa")
    r = montar()
    r["infraestrutura"]["scores"]["score_final"] = 1.3
    _falha(r, "Rodou com localização precisa")
    r = montar()
    r["infraestrutura"] = {"id_execucao": ID, "status": "endereco_nao_localizado"}
    r["preco_estimado"]["aviso_infraestrutura"] = None
    _falha(r, "Rodou com localização precisa")              # sem o aviso de nota neutra
    r["preco_estimado"]["aviso_infraestrutura"] = "Tempo estimado sem a infraestrutura..."
    _passa(r, "Rodou com localização precisa")


# ---------------------------------------------------------------- agente 5
def test_ag5_refaz_a_conta():
    _passa(montar(), "Refazer a conta de forma independente")
    _passa(montar(area_terreno=250, n_terrenos=6), "Refazer a conta de forma independente")   # separa o terreno
    r = montar()
    r["preco_estimado"]["avaliacao_planilha"]["valor_medio_imovel"] += 50
    _falha(r, "Refazer a conta de forma independente")


def test_ag5_refaz_a_conta_com_terreno_acentuado_e_preco_m2_do_anuncio():
    # Mesma regra do Agente 5 para o R$/m2 do terreno: topografia "acentuado" vale 80%
    # e, com area invalida, vale o preco/m2 do anuncio.
    r = montar(area_terreno=250, n_terrenos=6, extras_terreno={
        0: {"topografia": "Aclive acentuado"}, 1: {"area": 5, "pricePerSqm": 700}})
    usados_t = r["preco_estimado"]["comparaveis_usados"]["terreno"]
    assert len(usados_t) == 6 and r["preco_estimado"]["calculo_terreno"]["aplicado"]
    acentuado = next(t for t in usados_t if t["topografia"] == "Aclive acentuado")
    assert acentuado["conta_m2"][0]["valor_m2"] == round(200_000 / 300 * 0.8, 2)
    _passa(r, "Refazer a conta de forma independente")
    # Sem a regra (preco / area em todos), a conta daria outro valor: o teste pega o caso.
    sem_regra = copy.deepcopy(r)
    for t in sem_regra["preco_estimado"]["comparaveis_usados"]["terreno"]:
        t["topografia"], t["preco_m2_anuncio"] = None, None
        if t["area"] == 5:
            t["area"] = 1000
    _falha(sem_regra, "Refazer a conta de forma independente")


def test_ag5_refaz_a_conta_de_apartamento():
    aptos = [{"propertyType": "apartamento", "price": p, "area": a, "url": f"https://a/{p}"}
             for p, a in ((300000, 93), (499900, 120), (630000, 110), (648000, 117), (650000, 77), (790000, 106))]
    ag5 = executar_agente5({"propertyType": "apartamento", "area": 109}, [], aptos, {}, {})
    assert abs(refazer_conta_agente5(ag5) - ag5["avaliacao_planilha"]["valor_medio_imovel"]) <= 1


def test_ag5_usados_na_zona():
    _passa(montar(), "Todos os comparáveis usados estão na zona")
    r = montar()
    r["preco_estimado"]["comparaveis_usados"]["construcao"][0]["url"] = "https://longe/1"
    _falha(r, "Todos os comparáveis usados estão na zona")


def test_ag5_anuncio_do_proprio_imovel_fora_dos_usados():
    _passa(montar(), "Nenhum anúncio do próprio imóvel entrou nos usados")
    r = montar()
    r["comparaveis"][0]["eh_anuncio_do_alvo"] = True
    _falha(r, "Nenhum anúncio do próprio imóvel entrou nos usados")


def test_ag5_terrenos_coerentes_com_a_decisao():
    _passa(montar(), "Terrenos usados coerentes")                       # nao separou, 0 usados
    _passa(montar(area_terreno=250, n_terrenos=6), "Terrenos usados coerentes")   # separou, 6 usados
    r = montar()
    r["preco_estimado"]["comparaveis_usados"]["terreno"] = [{"url": "https://t/0", "preco": 200000, "area": 300}]
    v = _falha(r, "Terrenos usados coerentes")   # o bug de antes do commit 0b47594
    assert v["encontrado"] == "não separou o terreno; 1 terrenos usados", v


def test_ag1_repetidos_provaveis_entre_os_usados():
    _passa(montar(), "Repetidos prováveis entre os usados")
    r = montar()
    usados = r["preco_estimado"]["comparaveis_usados"]["construcao"]
    # O mesmo imovel em 3 anuncios: mesma rua (com e sem numero, grafia "R"), preco e area.
    for u, rua in zip(usados[:3], ("Rua Manuel Vaz, 12", "R Manuel Vaz", "Rua Manuel Vaz,")):
        u.update(rua=rua, preco=960_000, area=150)
    usados[1]["preco"] = 962_000          # 0,2%: ainda o mesmo preco
    v = _v(r, "Repetidos prováveis entre os usados")
    assert v["ok"] == "alerta", v                                      # aviso, nao falha
    assert v["encontrado"].startswith("1 grupo(s), 3 anúncios: Rua Manuel Vaz, 12 R$ 960.000 / 150 m² ×3"), v
    # Rua diferente, ou area 3% maior: nao junta.
    usados[2]["rua"] = "Rua Pedro Rendon"
    usados[1]["area"] = 155
    assert _v(r, "Repetidos prováveis entre os usados")["ok"] is True


def test_ag5_sem_valor_so_com_poucos_comparaveis():
    _passa(montar(), "Sem valor ⇔ menos comparáveis")
    _passa(montar(n_casas=4), "Sem valor ⇔ menos comparáveis")           # 4 < 5: sem valor, coerente
    r = montar()
    r["preco_estimado"]["avaliacao_confiavel"] = False                   # sem valor com 6 validos
    _falha(r, "Sem valor ⇔ menos comparáveis")


# ---------------------------------------------------------------- tela e Excel
def _com_uma_falha():
    r = montar()
    r["infraestrutura"]["id_execucao"] = "rodada-anterior"
    r["status"] = "completo"
    r["verificacoes"] = verificar_execucao(r)
    return r


def test_conferencia_tudo_certo_em_frases_simples():
    from services.verificacoes import resumo_conferencia
    c = resumo_conferencia(montar())
    assert c["titulo"] == "Conferência: tudo certo, os 5 agentes funcionaram como esperado"
    assert c["frases"][0] == ("Agente 1: buscou 10 anúncios de Curitiba/PR, todos desta avaliação. "
                              "Juntou 1 anúncio repetido."), c["frases"]
    assert [f.split(":")[0] for f in c["frases"]] == ["Agente 1", "Agente 2", "Agente 3", "Agente 4", "Agente 5"]
    texto = " ".join(c["frases"])
    assert "esperado" not in texto and "encontrado" not in texto and "exec-1" not in texto


def test_conferencia_com_problema_diz_o_problema_em_uma_frase():
    from services.verificacoes import resumo_conferencia
    r = montar()
    r["comparaveis"][0]["city"] = "Colombo"
    r["preco_estimado"]["avaliacao_planilha"]["valor_medio_imovel"] += 5000
    c = resumo_conferencia(r)
    assert c["titulo"] == "Conferência: ⚠️ 2 problemas (Agente 1, Agente 5)", c["titulo"]
    assert "⚠️ Agente 1: 1 anúncio é de outra cidade (Colombo/PR)." in c["frases"]
    assert any(f.startswith("⚠️ Agente 5: a conta refeita dá R$") for f in c["frases"]), c["frases"]
    assert any(f.startswith("Agente 2: separou") for f in c["frases"])


def test_tela_mostra_a_conferencia_no_fim_fechada():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file("app/interface.py", default_timeout=60)
    at.session_state["resultado"] = _com_uma_falha()
    at.session_state["imovel_alvo_dados"] = {"area": 120, "tipo": "house"}
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert not [e for e in at.expander if e.label.startswith("🧪")]      # o bloco de cima saiu
    conf = next(e for e in at.expander if e.label.startswith("Conferência:"))
    assert conf.label == "Conferência: ⚠️ 1 problema (Geral)" and conf.proto.expanded is False, conf.label
    frases = next(m.value for m in conf.markdown if m.value.startswith("- "))
    assert frases.startswith("- ⚠️ Conferência geral: usou dados de outra avaliação (Agente 4).") and "esperado" not in frases
    # No fim: depois de "Como chegamos neste valor".
    rotulos = [e.label for e in at.expander]
    assert rotulos.index(conf.label) > rotulos.index("💰 Agente Estimador de Preço")


def test_excel_tem_a_aba_verificacoes():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado
    r = _com_uma_falha()
    livro = load_workbook(io.BytesIO(excel_resultado({"alertas": [], "verificacoes": r["verificacoes"]}, [])))
    linhas = list(livro["Verificações"].iter_rows(values_only=True))
    assert linhas[0] == ("Resultado", "Agente", "Verificação", "Esperado", "Encontrado")   # o detalhe tecnico continua
    assert linhas[1][0] == "❌ FALHOU" and linhas[-1][0] == "✅ OK" and len(linhas) == 1 + len(r["verificacoes"])
    assert "Verificações" not in load_workbook(io.BytesIO(excel_resultado({"alertas": []}, []))).sheetnames


def test_banco_grava_a_coluna():
    from services.banco import COLUNAS_EXECUCOES, COLUNAS_JSON
    assert "verificacoes" in COLUNAS_EXECUCOES and "verificacoes" in COLUNAS_JSON


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
        except Exception as e:
            falhas += 1
            print(f"ERRO  {nome}: {type(e).__name__}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
