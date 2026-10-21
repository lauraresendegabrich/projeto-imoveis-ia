"""
Tela: alertas terreno_fora_do_padrao, possivel_repetido_no_calculo e
area_diferente_entre_portais, com explicacao e os anuncios num expander; laudo com a
mesma frase apontando para a tela e o Excel. Simulador do Streamlit. Sem rede.

    .venv/Scripts/python.exe -m tests.test_tela_alertas
"""
import logging
import sys

sys.path.insert(0, ".")
logging.disable(logging.CRITICAL)

from streamlit.testing.v1 import AppTest

from services.alertas import calcular_alertas
from tests.test_tela_amostra_concentrada import _laudo


def comp(preco, area, rua, **extra):
    return {"preco": preco, "area": area, "rua": rua, "bairro": "Centro", "portal": "VivaReal",
            "url": f"https://x/{rua}", "status_zona": "confirmado", **extra}


CASAS = [comp(400_000 + i * 20_000, 100 + i * 7, f"Rua B, {i}") for i in range(6)]
CASOS = {
    "terreno": {"construcao": CASAS, "terreno": [comp(110_000 + i * 1_000, 200, f"Rua T, {i}") for i in range(8)]
                + [comp(11_200, 200, "Av. X, 505"), comp(950_000, 9_500, "Rua Y, 3")]},
    "repetido": {"construcao": CASAS + [comp(550_000, 142, "Rua Elvio, 314"), comp(550_000, 142, "Rua Projetada, 314")],
                 "terreno": []},
    "area": {"construcao": CASAS[:5] + [comp(590_000, 80, "Rua C, 1", areas_entre_portais=[
        {"portal": "VivaReal", "area": 80.0, "url": "https://vr/1"},
        {"portal": "Chaves na Mão", "area": 110.0, "url": "https://cnm/2"}])], "terreno": []},
    "nenhum": {"construcao": CASAS, "terreno": []},
}


def _tela(caso):
    ag5 = {"status_avaliacao": "ok", "avaliacao_confiavel": True, "comparaveis_usados": CASOS[caso],
           "avaliacao_planilha": {"valor_medio_imovel": 500000, "valor_liquidez": 450000},
           "liquidez_experimental": {"tempo_estimado": "90 a 150 dias"}}
    alertas = calcular_alertas({"status_agente5": "ok", "qtd_comparaveis_construcao": 10}, ag5)
    at = AppTest.from_file("app/interface.py", default_timeout=60)
    at.session_state["resultado"] = {"status": "completo", "preco_estimado": ag5, "resumo": {},
                                     "comparaveis": [], "terrenos": [],
                                     "registro": {"execucao": {"alertas": alertas}, "comparaveis": []}}
    at.session_state["imovel_alvo_dados"] = {"area": 80, "tipo": "house"}
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def _aviso(at, inicio):
    return [w.value for w in at.warning if inicio in w.value]


def _tabela_com(at, texto):
    return [m.value for m in at.markdown if texto in m.value and "| Endereço" in m.value]


def test_terreno_fora_do_padrao():
    at = _tela("terreno")
    aviso = _aviso(at, "Terreno fora do padrão: 2 dos 10 terrenos usados no cálculo")
    assert aviso and "valor típico de R\$ 562/m²" in aviso[0] and "conservador" not in aviso[0], aviso
    assert "Ver os anúncios" in [e.label for e in at.expander]
    tabela = _tabela_com(at, "Av. X, 505")
    assert tabela and "| Motivo" in tabela[0] and "muito abaixo da mediana" in tabela[0], tabela
    assert any(l.endswith("confira esses anúncios na tela ou no Excel.")
               for l in _laudo(at).splitlines() if "Terreno fora do padrão" in l)


def test_possivel_repetido_no_calculo():
    at = _tela("repetido")
    aviso = _aviso(at, "Possível repetido no cálculo: 2 anúncios, em 1 grupo,")
    assert aviso and "pesando mais de uma vez na média. Confira os anúncios." in aviso[0], aviso
    # Logo abaixo do aviso, agrupado: titulo do grupo e os anuncios, sem a coluna Motivo.
    assert "Ver os 2 anúncios que podem ser repetidos" in [e.label for e in at.expander]
    assert any(m.value == "**Grupo 1: nº 314, R\\$ 550.000, 142 m²**" for m in at.markdown), \
        [m.value for m in at.markdown if "Grupo" in m.value]
    tabela = _tabela_com(at, "Rua Projetada, 314")
    assert tabela and "[ver](https://x/Rua Elvio, 314)" in tabela[0] and "| Motivo" not in tabela[0], tabela
    assert any(l.endswith("Confira os anúncios na tela ou no Excel.")
               for l in _laudo(at).splitlines() if "Possível repetido" in l)


def test_area_diferente_entre_portais():
    at = _tela("area")
    aviso = _aviso(at, "Área diferente entre portais: 1 imóvel usado no cálculo aparece")
    assert aviso and "Foi usada a menor área" in aviso[0], aviso
    tabela = _tabela_com(at, "Rua C, 1")
    assert tabela and "VivaReal 80 m² · Chaves na Mão 110 m²; usada 80 m²" in tabela[0], tabela
    assert any(l.endswith("Confira os anúncios na tela ou no Excel.")
               for l in _laudo(at).splitlines() if "Área diferente entre portais" in l)


def test_nenhum_alerta_nada_na_tela():
    at = _tela("nenhum")
    textos = " ".join(w.value for w in at.warning)
    for inicio in ("Terreno fora do padrão", "Possível repetido", "Área diferente entre portais",
                   "Amostra concentrada"):
        assert inicio not in textos, inicio
    assert not [e.label for e in at.expander if e.label.startswith("Ver os ")]


def test_excel_usa_os_mesmos_textos():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado
    for caso, inicio in (("terreno", "Terreno fora do padrão:"), ("repetido", "Possível repetido no cálculo:"),
                         ("area", "Área diferente entre portais:")):
        alertas = calcular_alertas({"status_agente5": "ok", "qtd_comparaveis_construcao": 10},
                                   {"comparaveis_usados": CASOS[caso]})
        motivos = [l[1] for l in load_workbook(io.BytesIO(excel_resultado({"alertas": alertas}, [])))["Alertas"]
                   .iter_rows(values_only=True) if l[1]]
        assert any(str(m).startswith(inicio) for m in motivos), (caso, motivos[:3])
        assert not any("conservador" in str(m) for m in motivos)


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
