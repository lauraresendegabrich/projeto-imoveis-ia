"""
Funil dos anuncios (app/funil.py): em que agente cada anuncio saiu e por que.
Sem rede.

    .venv/Scripts/python.exe -m tests.test_funil
"""
import logging
import os
import sys

sys.path.insert(0, ".")
logging.disable(logging.CRITICAL)

from app.funil import (ENTROU, SAIU_AG1, SAIU_AG2_PARECIDOS, SAIU_AG2_ZONA, SAIU_AG5, funil_ligado,
                       montar_funil, resumo_funil)


def ad(n, cluster="A", **extra):
    return {"url": f"https://x/{n}", "price": 500000 + n, "area": 100, "bedrooms": 3, "neighborhood": "Centro",
            "source": "VivaReal", "cluster": cluster, **extra}


RESULTADO = {
    "comparaveis": [
        ad(1),                                                    # entrou
        ad(2, eh_anuncio_do_alvo=True),                           # proprio anuncio
        ad(3),                                                    # descartado no Ag5
        ad(4, distancia_metros=900),                              # fora da zona
        ad(5),                                                    # nao verificado
        ad(6, "B", pre_classificacao={"elegivel_llm": False, "motivos_incompatibilidade": [
            "area construida difere 66.7% do alvo (120.0m² vs 360.0m²; limite 30%)"]}),
        ad(7, "B", pre_classificacao={"elegivel_llm": True}, justificativa="Área 35% menor e sem piscina."),
        ad(8, origem_busca="proximidade", distancia_alvo_m=700),  # entrou, veio por proximidade
    ],
    "terrenos": [ad(9, "terreno", propertyType="Terrenos")],
    "zona_homogenea": {"coordenadas_alvo": {"lat": -27.8, "lon": -48.6}, "raio_final_m": 450,
                       "comparaveis_confirmados": [ad(1), ad(2), ad(3), ad(8)],
                       "comparaveis_nao_verificados": [ad(5)], "terrenos": [ad(9)]},
    "preco_estimado": {
        "comparaveis_usados": {"construcao": [
            {"url": "https://x/1", "status_zona": "confirmado", "conta_m2": [{"valor_m2": 5000.0, "na_media": True}]},
            {"url": "https://x/8", "status_zona": "confirmado", "conta_m2": [{"valor_m2": 9100.4, "na_media": False}]},
        ], "terreno": []},
        "comparaveis_descartados": {
            "construcao": [{"url": "https://x/3", "motivo_descarte": "R$/m² da construção fora da faixa plausível"}],
            "terreno": [{"url": "https://x/9", "motivo_descarte": "terreno não separado: só 1 terreno(s) válidos"}]},
    },
    "duplicatas_descartadas": [{"Mantido": "VivaReal", "Preço": 500001, "Área": 100, "Descartado": "ImovelWeb",
                                "Anúncio descartado": "https://iw/1", "Motivo": "mesmo_codigo_e_anunciante"}],
}


def _por_url(linhas):
    return {l["Anúncio"]: l for l in linhas}


def test_cada_anuncio_tem_etapa_e_motivo():
    l = _por_url(montar_funil(RESULTADO))
    assert l["https://x/1"]["Etapa"] == ENTROU == "✅ Chegou ao cálculo"
    assert l["https://x/1"]["Motivo"] == "R$ 5.000/m², usado na média"
    assert l["https://x/1"]["Zona"] == "confirmado na zona"
    assert l["https://x/8"]["Motivo"] == "R$ 9.100/m², cortado pela média aparada"
    assert l["https://x/2"]["Etapa"] == SAIU_AG5 and "é o anúncio do próprio imóvel" in l["https://x/2"]["Motivo"]
    assert l["https://x/3"]["Etapa"] == SAIU_AG5 and "fora da faixa plausível" in l["https://x/3"]["Motivo"]
    assert l["https://x/4"]["Etapa"] == SAIU_AG2_ZONA and "está a 900 m" in l["https://x/4"]["Motivo"]
    assert "raio usado foi de 450 m" in l["https://x/4"]["Motivo"]
    assert l["https://x/5"]["Etapa"] == SAIU_AG2_ZONA and "posição não confirmada" in l["https://x/5"]["Motivo"]
    assert l["https://x/6"]["Etapa"] == SAIU_AG2_PARECIDOS and l["https://x/6"]["Motivo"].startswith("pré-classificação: area construida difere 66.7%")
    assert l["https://x/7"]["Etapa"] == SAIU_AG2_PARECIDOS and "a IA não considerou parecido: Área 35% menor" in l["https://x/7"]["Motivo"]
    assert l["https://x/8"]["Etapa"] == ENTROU and l["https://x/8"]["Veio por"] == "proximidade (700 m)"
    assert l["https://x/9"]["Etapa"] == SAIU_AG5 and l["https://x/9"]["Tipo"] == "Terreno"
    assert l["https://iw/1"]["Etapa"] == SAIU_AG1 and "repetido" in l["https://iw/1"]["Motivo"]


def test_nomes_das_etapas():
    assert (SAIU_AG1, SAIU_AG2_PARECIDOS, SAIU_AG2_ZONA, SAIU_AG5) == (
        "Saiu no Agente 1 (busca)", "Saiu no Agente 2 (parecidos)", "Saiu no Agente 2 (zona)",
        "Saiu no Agente 5 (cálculo)")


def test_motivo_na_conta():
    from app.funil import motivo_na_conta
    ok = {"status_avaliacao": "ok"}
    copia = {"conta_m2": [{"valor_m2": 5238.1, "na_media": True}, {"valor_m2": 5238.1, "na_media": False}]}
    assert motivo_na_conta(copia, ok) == ("R$ 5.238/m², usado na média uma vez (a outra cópia foi cortada "
                                          "pela média aparada)")
    duas = {"conta_m2": [{"valor_m2": 4100, "na_media": True}, {"valor_m2": 4800, "na_media": False}]}
    assert motivo_na_conta(duas, ok) == "R$ 4.100/m², usado na média; R$ 4.800/m², cortado pela média aparada"
    sem = {"status_avaliacao": "amostra_insuficiente", "comparaveis_usados": {"construcao": [{}, {}, {}, {}]}}
    assert motivo_na_conta(duas, sem) == "não houve cálculo: 4 comparáveis < mínimo de 5"
    assert motivo_na_conta({}, ok) == "usado no cálculo"


def test_conta_vem_do_agente5():
    # Os R$/m2 de cada anuncio e o corte da media aparada batem com o valor do Agente 5.
    from agents.price_liquidity import executar_agente5
    comps = [{"propertyType": "apartamento", "price": p, "area": a, "url": f"https://a/{p}"}
             for p, a in ((300000, 93), (499900, 120), (630000, 110), (648000, 117), (650000, 77), (790000, 106))]
    r = executar_agente5({"propertyType": "apartamento", "area": 109}, [], comps, {}, {})
    usados = r["comparaveis_usados"]["construcao"]
    na_media = [c["valor_m2"] for u in usados for c in u["conta_m2"] if c["na_media"]]
    assert round(sum(na_media) / len(na_media), 2) == r["calculo_construcao"]["valor_m2_referencia"] == 5721.1
    cortados = {u["url"] for u in usados if not u["conta_m2"][0]["na_media"]}
    assert cortados == {"https://a/300000", "https://a/650000"}, cortados   # o mais barato e o mais caro


def test_resumo_e_ordem():
    linhas = montar_funil(RESULTADO)
    assert resumo_funil(linhas) == {ENTROU: 2, SAIU_AG5: 3, SAIU_AG2_ZONA: 2, SAIU_AG2_PARECIDOS: 2, SAIU_AG1: 1}
    assert linhas[0]["Etapa"] == ENTROU and linhas[-1]["Etapa"] == SAIU_AG1


def test_imovel_nao_localizado():
    r = dict(RESULTADO, zona_homogenea={"status": "zona_nao_verificada", "comparaveis_confirmados": []},
             preco_estimado={})
    l = _por_url(montar_funil(r))
    assert l["https://x/4"]["Etapa"] == SAIU_AG2_ZONA and "não localizado no mapa" in l["https://x/4"]["Motivo"]


def test_liga_e_desliga():
    original = os.environ.get("MOSTRAR_FUNIL")
    try:
        os.environ.pop("MOSTRAR_FUNIL", None)
        assert funil_ligado()
        os.environ["MOSTRAR_FUNIL"] = "0"
        assert not funil_ligado()
    finally:
        if original is None:
            os.environ.pop("MOSTRAR_FUNIL", None)
        else:
            os.environ["MOSTRAR_FUNIL"] = original


def test_excel_tem_a_aba_funil():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado
    livro = load_workbook(io.BytesIO(excel_resultado({"alertas": []}, [], funil=montar_funil(RESULTADO))))
    linhas = list(livro["Funil"].iter_rows(values_only=True))
    assert linhas[0][:3] == ("Etapa", "Motivo", "Zona") and len(linhas) == 1 + 10
    assert "R$ 5.000/m², usado na média" in [l[1] for l in linhas]
    assert "Funil" not in load_workbook(io.BytesIO(excel_resultado({"alertas": []}, []))).sheetnames


def test_tela_mostra_o_funil():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file("app/interface.py", default_timeout=60)
    preco = {"status_avaliacao": "ok", "avaliacao_confiavel": True,
             "avaliacao_planilha": {"valor_medio_imovel": 500000, "valor_liquidez": 450000},
             "liquidez_experimental": {"tempo_estimado": "90 a 150 dias"}, **RESULTADO["preco_estimado"]}
    at.session_state["resultado"] = {**RESULTADO, "status": "completo", "preco_estimado": preco, "resumo": {}}
    at.session_state["imovel_alvo_dados"] = {"area": 100, "tipo": "house"}
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    assert any(e.label.startswith("🔎 Funil dos anúncios (conferência) — 10 anúncios") for e in at.expander)
    resumo = [m.value for m in at.markdown if "Chegou ao cálculo" in m.value and "Saiu no Agente 1" in m.value]
    assert resumo, [m.value for m in at.markdown][-5:]
    tabela = at.dataframe[-1].value
    assert list(tabela.columns[:2]) == ["Etapa", "Motivo"] and len(tabela) == 10
    # "–" no lugar de None e preco com ponto de milhar.
    assert "None" not in tabela.drop(columns=["Anúncio"]).astype(str).to_string()
    assert "R$ 500.001" in set(tabela["Preço"]) and "–" in set(tabela["Quartos"])


def test_tela_formata_vazio_e_preco():
    from app.funil import linhas_para_tela
    linha = linhas_para_tela([{"Etapa": "x", "Preço": 1200000.0, "Área (m²)": 179.5, "Quartos": None,
                               "Nota vizinhança (Ag3)": 0.71, "Anúncio": None}])[0]
    assert linha == {"Etapa": "x", "Preço": "R$ 1.200.000", "Área (m²)": "179,5", "Quartos": "–",
                     "Nota vizinhança (Ag3)": "0,71", "Anúncio": None}, linha
    assert linhas_para_tela([{"Área (m²)": 360.0}])[0]["Área (m²)"] == "360"


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
