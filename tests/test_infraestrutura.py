"""
Testes do Agente 4 sem localizacao precisa e do tempo de venda sem nota de
infraestrutura (Agente 5). Sem rede.

    .venv/Scripts/python.exe -m tests.test_infraestrutura
"""
import sys

sys.path.insert(0, ".")

import agents.comparables as ag2
import agents.infra_evaluator as ag4
from agents.price_liquidity import AVISO_FALHA_INFRAESTRUTURA, AVISO_SEM_INFRAESTRUTURA, executar_agente5

ALVO = {"propertyType": "apartamento", "area": 80}
COMPS = [{"propertyType": "apartamento", "price": 400000 + i * 10000, "area": 80} for i in range(5)]


def _ag5(dados_ag4):
    return executar_agente5(ALVO, [], COMPS, {}, dados_ag4)


def test_com_infraestrutura_como_hoje():
    r = _ag5({"scores": {"score_final": 0.8}})
    liq = r["liquidez_experimental"]
    # 0,35 x 0,50 (qualidade neutra) + 0,40 x 0,80 + 0,25 x 0,90 = 0,72
    assert liq["score_liquidez"] == 0.72 and liq["tempo_estimado"] == "60 a 90 dias", liq
    assert r["aviso_infraestrutura"] is None and liq["aviso_infraestrutura"] is None


NAO_LOCALIZADO = {"status": "endereco_nao_localizado", "scores": {},
                  "mensagem": "Infraestrutura não avaliada (endereço não localizado)"}


def test_sem_infraestrutura_usa_neutro_e_avisa():
    r = _ag5(NAO_LOCALIZADO)
    liq = r["liquidez_experimental"]
    # 0,35 x 0,50 + 0,40 x 0,50 (infraestrutura neutra) + 0,25 x 0,90 = 0,60 (pesos de sempre)
    assert liq["score_liquidez"] == 0.6 and liq["tempo_estimado"] == "90 a 150 dias", liq
    assert liq["score_agente4_usado"] is None
    assert r["aviso_infraestrutura"] == AVISO_SEM_INFRAESTRUTURA == (
        "Tempo estimado sem a infraestrutura do entorno (endereço não localizado no mapa): a "
        "infraestrutura foi considerada média. Confira o entorno do imóvel.")


def test_falha_do_agente4_tem_outro_aviso():
    # Sem nota e sem o status de "nao localizado" (Agente 4 falhou ou nao gravou nada).
    r = _ag5({})
    assert r["liquidez_experimental"]["score_liquidez"] == 0.6           # mesmo neutro 0,50
    assert r["aviso_infraestrutura"] == AVISO_FALHA_INFRAESTRUTURA == (
        "Tempo estimado sem a infraestrutura do entorno (a avaliação da infraestrutura falhou "
        "nesta rodada): a infraestrutura foi considerada média.")


def test_agente4_nao_localizado_nao_usa_o_centro_da_cidade():
    chamadas = []

    def proibido(*a, **k):
        chamadas.append(a)
        raise AssertionError("nao pode usar a geocodificacao propria nem buscar pontos de interesse")

    import json, os, tempfile
    originais = (ag2._geocodificar_com_fallback, ag4._geocodificar, ag4._buscar_pois_classificados, ag4.DATA_DIR)
    ag2._geocodificar_com_fallback = lambda **k: (None, None, None, None)
    ag4._geocodificar = proibido                 # antes: aceitava o centro da cidade
    ag4._buscar_pois_classificados = proibido
    try:
        with tempfile.TemporaryDirectory() as pasta:
            ag4.DATA_DIR = pasta                 # o Agente 4 grava o status; nao suja data/
            r = ag4.avaliar_infraestrutura(
                imovel_alvo={"rua": "Rua Inexistente", "numero": "", "bairro": "Bairro W",
                             "cidade": "Itaboraí", "estado": "RJ"},
                id_execucao="teste-infra-sem-localizacao",
            )
            gravado = json.load(open(os.path.join(pasta, "infra_avaliada_ag4.json"), encoding="utf-8"))
    finally:
        (ag2._geocodificar_com_fallback, ag4._geocodificar, ag4._buscar_pois_classificados,
         ag4.DATA_DIR) = originais
    assert chamadas == []
    # Gravado com o id da avaliacao: o Agente 5 le e usa o aviso de "nao localizado".
    assert gravado["status"] == "endereco_nao_localizado" and gravado["id_execucao"] == "teste-infra-sem-localizacao"
    assert r["status"] == "endereco_nao_localizado" and r["scores"] == {}
    assert r["mensagem"] == "Infraestrutura não avaliada (endereço não localizado)"


def test_banco_e_excel_com_aviso_de_infraestrutura():
    import io
    from openpyxl import load_workbook
    from services.registro_execucao import excel_resultado, montar_registro
    r = _ag5(NAO_LOCALIZADO)
    e, comps = montar_registro(id_execucao="x", imovel_alvo={}, leilao_entrada={}, leilao={},
                               imoveis_coletados=[], resultado_ag2={}, zona_resultado={}, resultado_ag3={},
                               resultado_ag4={}, resultado_ag5=r)
    assert e["parametros"]["aviso_infraestrutura"] == AVISO_SEM_INFRAESTRUTURA
    resumo = {l[0]: l[1] for l in load_workbook(io.BytesIO(excel_resultado(e, comps)))["Resumo"].iter_rows(values_only=True)}
    assert resumo["Tempo de venda sem infraestrutura"].startswith("Tempo estimado sem a infraestrutura do entorno")


if __name__ == "__main__":
    testes = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    falhas = 0
    for nome, fn in testes:
        try:
            fn()
            print(f"OK    {nome}")
        except AssertionError as e:
            falhas += 1
            print(f"FALHA {nome}: {e}")
    print(f"\n{len(testes) - falhas}/{len(testes)} testes passaram")
    sys.exit(1 if falhas else 0)
