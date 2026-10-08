"""
Monta o registro de uma avaliacao para o banco (tabelas execucoes e comparaveis,
ver db/schema.sql) a partir do que a avaliacao ja produziu: formulario, Agentes
1 a 5 e avaliacao de leilao. Nao faz I/O de banco (isso e services/banco.py).
"""
from __future__ import annotations

import os
from collections import Counter
from pathlib import Path

from services.leilao import FATOR_LUCRO_LEILAO, ORIGEM_CAIXA

RAIZ = Path(__file__).resolve().parent.parent

TIPOS_ALVO = {"house": "casa", "apartment": "apartamento"}


def versao_codigo(raiz: Path = RAIZ) -> str:
    """Hash do commit atual lendo .git diretamente (o Streamlit Cloud pode nao ter o git)."""
    try:
        git = raiz / ".git"
        head = (git / "HEAD").read_text(encoding="utf-8").strip()
        if not head.startswith("ref:"):
            return head[:40]
        ref = head.split(" ", 1)[1].strip()
        arquivo_ref = git / ref
        if arquivo_ref.exists():
            return arquivo_ref.read_text(encoding="utf-8").strip()[:40]
        empacotados = git / "packed-refs"
        if empacotados.exists():
            for linha in empacotados.read_text(encoding="utf-8").splitlines():
                if linha.endswith(" " + ref):
                    return linha.split(" ", 1)[0][:40]
    except Exception:
        pass
    return "desconhecida"


def _num(valor):
    try:
        numero = float(valor)
    except (TypeError, ValueError):
        return None
    return numero if numero == numero else None


def _int(valor):
    numero = _num(valor)
    return int(numero) if numero is not None else None


def _modelos_ia(resultado_ag2, zona_resultado, resultado_ag3, resultado_ag4) -> dict:
    """Servico de IA usado em cada agente, com quantas chamadas cada um resolveu."""
    ag2 = Counter(
        f"{c.get('provider_llm')}/{c.get('modelo_llm')}"
        for c in (resultado_ag2 or {}).get("comparaveis", [])
        if c.get("status_julgamento") == "JULGADO_LLM"
    )
    ag3_analises = [((resultado_ag3 or {}).get("imovel_alvo") or {}).get("analise_qualitativa") or {}]
    ag3_analises += [c.get("analise_qualitativa") or {} for c in (resultado_ag3 or {}).get("comparaveis", [])]
    ag3 = Counter(a.get("llm_usada") for a in ag3_analises if a.get("llm_usada"))
    return {
        "agente2_classificacao": dict(ag2),
        "agente2_zona_visao": ((zona_resultado or {}).get("zona_homogenea") or {}).get("provider_visao"),
        "agente3": dict(ag3),
        "agente4": ((resultado_ag4 or {}).get("scores") or {}).get("provider_llm"),
    }


def _linhas_comparaveis(id_execucao: str, resultado_ag5: dict) -> list[dict]:
    linhas = []
    for chave, entrou in (("comparaveis_usados", True), ("comparaveis_descartados", False)):
        grupos = (resultado_ag5 or {}).get(chave) or {}
        for tipo in ("construcao", "terreno"):
            for c in grupos.get(tipo) or []:
                preco, area = _num(c.get("preco")), _num(c.get("area"))
                linhas.append({
                    "id_execucao": id_execucao,
                    "tipo": tipo,
                    "preco": preco,
                    "area": area,
                    "valor_m2": round(preco / area, 2) if preco and area else None,
                    "endereco": c.get("rua") or None,
                    "bairro": c.get("bairro") or None,
                    "portal": c.get("portal") or None,
                    "link": c.get("url") or None,
                    "status_zona": c.get("status_zona") or "confirmado",
                    "entrou_no_calculo": entrou,
                })
    return linhas


def montar_registro(
    *,
    id_execucao: str,
    imovel_alvo: dict,
    leilao_entrada: dict | None,
    leilao: dict | None,
    imoveis_coletados: list | None,
    resultado_ag2: dict | None,
    zona_resultado: dict | None,
    resultado_ag3: dict | None,
    resultado_ag4: dict | None,
    resultado_ag5: dict | None,
    tempos: dict | None = None,
    falhas: list | None = None,
) -> tuple[dict, list[dict]]:
    """Retorna (linha de execucoes, linhas de comparaveis)."""
    from agents.price_liquidity import MIN_AMOSTRAS_CONFIAVEL
    from agents.comparables import MIN_CONFIRMADOS_ZONA

    imovel_alvo = imovel_alvo or {}
    leilao_entrada = leilao_entrada or {}
    leilao = leilao or {}
    ag5 = resultado_ag5 or {}
    zona = zona_resultado or {}
    resumo_zona = zona.get("resumo_zona") or {}
    raio = _int(resumo_zona.get("raio_usado_metros") or (zona.get("zona_homogenea") or {}).get("raio_metros"))
    usados = ag5.get("comparaveis_usados") or {}
    usados_todos = (usados.get("construcao") or []) + (usados.get("terreno") or [])
    avaliacao = ag5.get("avaliacao_planilha") or {}
    m2 = ag5.get("valor_m2_zona_homogenea") or {}
    confiavel = bool(ag5.get("avaliacao_confiavel"))
    coletados = imoveis_coletados or []

    try:
        from services.athena_client import _CACHE_COLETAS
        coletas = _CACHE_COLETAS.get("coletas")
    except Exception:
        coletas = None

    numero = str(imovel_alvo.get("numero") or "").strip()
    rua = str(imovel_alvo.get("rua") or imovel_alvo.get("street") or "").strip()

    execucao = {
        "id_execucao": id_execucao,
        "versao_codigo": versao_codigo(),
        "parametros": {
            "fator_lucro_leilao": FATOR_LUCRO_LEILAO,
            "min_amostras_confiavel": MIN_AMOSTRAS_CONFIAVEL,
            "min_confirmados_zona": MIN_CONFIRMADOS_ZONA,
            "raio_zona_m": raio,
        },
        "fontes": {
            "athena_coletas": coletas,
            "athena_tabela": os.getenv("ATHENA_TABLE") or "anuncios",
            "anuncios_apify": sum(1 for im in coletados if not im.get("fonte_dados")),
            "google_maps": bool(os.getenv("GOOGLE_MAPS_KEY")),
            "zona_calculada": ag5.get("zona_calculada_nesta_avaliacao"),
            "qwen_configurado": bool(os.getenv("QWEN_API_URL")),
        },
        "modelos_ia": _modelos_ia(resultado_ag2, zona_resultado, resultado_ag3, resultado_ag4),
        "tempo_s": (tempos or {}).get("total"),
        "falhas": {"falhas": list(falhas or []), "avisos_agente5": ag5.get("avisos") or [], "tempos": tempos or {}},

        "origem": leilao_entrada.get("origem") or ORIGEM_CAIXA,
        "codigo_caixa": leilao_entrada.get("codigo_caixa"),
        "link": leilao_entrada.get("link"),
        "modalidade": leilao_entrada.get("modalidade"),
        "tipo": TIPOS_ALVO.get(imovel_alvo.get("tipo"), imovel_alvo.get("tipo")),
        "endereco": f"{rua}, {numero}" if rua and numero else (rua or None),
        "bairro": imovel_alvo.get("bairro") or imovel_alvo.get("neighborhood"),
        "cidade": imovel_alvo.get("cidade") or imovel_alvo.get("city"),
        "uf": (imovel_alvo.get("estado") or imovel_alvo.get("state") or "")[:2] or None,
        "area_privativa": _num(imovel_alvo.get("area")),
        "area_terreno": _num(imovel_alvo.get("area_terreno")),
        "quartos": _int(imovel_alvo.get("bedrooms")),
        "banheiros": _int(imovel_alvo.get("bathrooms")),
        # O formulario exige quartos/area; nada e presumido por enquanto.
        "campos_presumidos": [],
        "valor_minimo": _num(leilao_entrada.get("valor_minimo")),
        "valor_avaliacao": _num(leilao_entrada.get("valor_avaliacao")),

        "anuncios_encontrados": len(coletados),
        "repetidos": sum(len(im.get("duplicatas") or []) for im in coletados),
        "parecidos": _int(((resultado_ag2 or {}).get("resumo") or {}).get("cluster_a")),
        "confirmados_na_zona": _int(resumo_zona.get("confirmados")),
        "usados_por_fallback": sum(1 for c in usados_todos if c.get("status_zona") in ("fallback", "sem_validacao")),
        "terrenos_usados": len(usados.get("terreno") or []),
        "raio_zona_m": raio,

        "valor_m2_construcao": _num(((m2.get("construcao") or {}).get("combinados") or {}).get("valor_m2_referencia")),
        # Sem terreno no calculo (ex.: apartamento) fica vazio, nao 0.
        "valor_m2_terreno": (
            _num((m2.get("terreno") or {}).get("valor_m2_referencia")) if usados.get("terreno") else None
        ),
        "valor_mercado": _num(avaliacao.get("valor_medio_imovel")),
        "liquidez": _num(avaliacao.get("valor_liquidez")),

        "status_agente5": ag5.get("status_avaliacao") or "sem_resultado",
        "qtd_comparaveis_construcao": len(usados.get("construcao") or []),
        "qtd_terrenos": len(usados.get("terreno") or []),
        "lance_maximo": _num(leilao.get("lance_maximo")),
        "sobra": _num(leilao.get("sobra")),
        "decisao": leilao.get("decisao"),
        "tempo_venda": (ag5.get("liquidez_experimental") or {}).get("tempo_estimado") if confiavel else None,
    }
    return execucao, _linhas_comparaveis(id_execucao, ag5)


# =============================================================================
# Excel do resultado da avaliacao (botao de exportacao na interface)
# =============================================================================
# Mesmo conteudo gravado no banco, em linguagem de usuario. Aba "Resumo" no formato
# Campo | Valor (na ordem do laudo) e aba "Comparaveis" com um anuncio por linha.

FORMATO_REAIS = '"R$" #,##0.00'
FORMATO_AREA = '#,##0.00 "m²"'
FORMATO_INTEIRO = "#,##0"

# (rotulo, chave em execucao, formato). Chave None = titulo de secao.
LINHAS_RESUMO = (
    ("Imóvel", None, None),
    ("Tipo", "tipo", None),
    ("Endereço", "endereco", None),
    ("Bairro", "bairro", None),
    ("Cidade", "cidade", None),
    ("UF", "uf", None),
    ("Área privativa", "area_privativa", FORMATO_AREA),
    ("Área do terreno", "area_terreno", FORMATO_AREA),
    ("Quartos", "quartos", FORMATO_INTEIRO),
    ("Banheiros", "banheiros", FORMATO_INTEIRO),
    ("Avaliação", None, None),
    ("Valor de mercado", "valor_mercado", FORMATO_REAIS),
    ("Valor de liquidez (−10%)", "liquidez", FORMATO_REAIS),
    ("Valor do m² da construção", "valor_m2_construcao", FORMATO_REAIS),
    ("Valor do m² do terreno", "valor_m2_terreno", FORMATO_REAIS),
    ("Tempo estimado de venda", "tempo_venda", None),
    ("Estimativa confiável", "_confiavel", None),
    ("Leilão CAIXA", None, None),
    ("Decisão", "decisao", None),
    ("Lance máximo", "lance_maximo", FORMATO_REAIS),
    ("Valor mínimo CAIXA", "valor_minimo", FORMATO_REAIS),
    ("Sobra (negativo = falta)", "sobra", FORMATO_REAIS),
    ("Valor de avaliação CAIXA (informativo)", "valor_avaliacao", FORMATO_REAIS),
    ("Modalidade", "modalidade", None),
    ("Código do imóvel na CAIXA", "codigo_caixa", None),
    ("Link do imóvel na CAIXA", "link", None),
    ("Como chegamos ao valor", None, None),
    ("Anúncios encontrados na região", "anuncios_encontrados", FORMATO_INTEIRO),
    ("Anúncios repetidos descartados", "repetidos", FORMATO_INTEIRO),
    ("Imóveis parecidos com o seu", "parecidos", FORMATO_INTEIRO),
    ("Confirmados na zona homogênea", "confirmados_na_zona", FORMATO_INTEIRO),
    ("Usados de fora da zona confirmada", "usados_por_fallback", FORMATO_INTEIRO),
    ("Comparáveis de construção no cálculo", "qtd_comparaveis_construcao", FORMATO_INTEIRO),
    ("Terrenos no cálculo", "qtd_terrenos", FORMATO_INTEIRO),
    ("Raio da zona homogênea (m)", "raio_zona_m", FORMATO_INTEIRO),
    ("Registro", None, None),
    ("Identificador da avaliação", "id_execucao", None),
)

ROTULOS_TIPO_COMPARAVEL = {"construcao": "Construção", "terreno": "Terreno"}
ROTULOS_STATUS_ZONA = {
    "confirmado": "Confirmado na zona",
    "fallback": "Fora da zona (fallback)",
    "sem_validacao": "Sem validação de distância",
}


def excel_resultado(execucao: dict, comparaveis: list[dict]) -> bytes:
    """Arquivo .xlsx com as abas Resumo e Comparaveis."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    negrito = Font(bold=True)
    titulo_secao = Font(bold=True, color="FFFFFF")
    fundo_secao = PatternFill("solid", fgColor="1F6F5C")
    fundo_cabecalho = PatternFill("solid", fgColor="E2EFE9")
    link = Font(color="0563C1", underline="single")

    livro = Workbook()

    # ── Resumo ────────────────────────────────────────────────────────────
    resumo = livro.active
    resumo.title = "Resumo"
    valores = dict(execucao)
    valores["_confiavel"] = "Sim" if execucao.get("status_agente5") == "ok" else "Não"
    for rotulo, chave, formato in LINHAS_RESUMO:
        if chave is None:
            resumo.append([rotulo, None])
            for celula in resumo[resumo.max_row]:
                celula.font, celula.fill = titulo_secao, fundo_secao
            continue
        valor = valores.get(chave)
        resumo.append([rotulo, valor])
        celula = resumo.cell(row=resumo.max_row, column=2)
        if formato and isinstance(valor, (int, float)):
            celula.number_format = formato
        if chave == "link" and valor:
            celula.hyperlink, celula.font = valor, link
        if chave == "decisao" and valor:
            celula.font = negrito
    resumo.column_dimensions["A"].width = 40
    resumo.column_dimensions["B"].width = 42
    for linha in resumo.iter_rows(min_col=2, max_col=2):
        linha[0].alignment = Alignment(horizontal="left")

    # ── Comparaveis ───────────────────────────────────────────────────────
    aba = livro.create_sheet("Comparáveis")
    colunas = (
        ("Tipo", 12, None), ("Entrou no cálculo", 17, None), ("Localização", 26, None),
        ("Preço", 16, FORMATO_REAIS), ("Área", 12, FORMATO_AREA), ("Valor do m²", 15, FORMATO_REAIS),
        ("Endereço", 36, None), ("Bairro", 22, None), ("Portal", 15, None), ("Anúncio", 60, None),
    )
    aba.append([nome for nome, _, _ in colunas])
    for celula in aba[1]:
        celula.font, celula.fill = negrito, fundo_cabecalho
    ordenados = sorted(
        comparaveis,
        key=lambda c: (c.get("tipo") or "", not c.get("entrou_no_calculo"), c.get("valor_m2") or 0),
    )
    for c in ordenados:
        aba.append([
            ROTULOS_TIPO_COMPARAVEL.get(c.get("tipo"), c.get("tipo")),
            "Sim" if c.get("entrou_no_calculo") else "Não",
            ROTULOS_STATUS_ZONA.get(c.get("status_zona"), c.get("status_zona")),
            c.get("preco"), c.get("area"), c.get("valor_m2"),
            c.get("endereco"), c.get("bairro"), c.get("portal"), c.get("link"),
        ])
        linha = aba.max_row
        for indice, (_, _, formato) in enumerate(colunas, start=1):
            if formato:
                aba.cell(row=linha, column=indice).number_format = formato
        if c.get("link"):
            celula = aba.cell(row=linha, column=len(colunas))
            celula.hyperlink, celula.font = c["link"], link
    for indice, (_, largura, _) in enumerate(colunas, start=1):
        aba.column_dimensions[aba.cell(row=1, column=indice).column_letter].width = largura
    aba.freeze_panes = "A2"
    if aba.max_row > 1:
        aba.auto_filter.ref = aba.dimensions

    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()
